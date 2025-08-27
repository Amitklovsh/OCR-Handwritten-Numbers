from __future__ import annotations

from dataclasses import dataclass
from typing import Dict, List

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class SheetInfo:
    sheet_name: str
    meta: Dict[str, object]


def _write_meta(ws: Worksheet, meta: Dict[str, object]) -> int:
    row = 1
    ws.cell(row=row, column=1, value="Metadata"); row += 1
    for key, value in meta.items():
        ws.cell(row=row, column=1, value=str(key))
        ws.cell(row=row, column=2, value=str(value))
        row += 1
    return row + 1  # leave a gap


def export_tables_to_workbook(output_path: str, tables: List[List[List[str]]], sheet_infos: List[SheetInfo]) -> None:
    wb = openpyxl.Workbook()
    # Remove default sheet
    default = wb.active
    wb.remove(default)

    manifest_rows: List[List[str]] = [["Sheet", "Image", "Engine", "Variant", "Rows", "Cols", "TotalMs", "Preprocess", "EngineOptions"]]

    for table, info in zip(tables, sheet_infos):
        ws = wb.create_sheet(title=info.sheet_name[:31])
        start_row = _write_meta(ws, info.meta)
        # Write table
        for r, row_vals in enumerate(table, start=start_row):
            for c, v in enumerate(row_vals, start=1):
                ws.cell(row=r, column=c, value=v)
        # Add manifest entry
        manifest_rows.append([
            info.sheet_name,
            str(info.meta.get("image_name")),
            str(info.meta.get("engine")),
            str(info.meta.get("variant_name")),
            str(len(table)),
            str(len(table[0]) if table else 0),
            str(info.meta.get("duration_ms")),
            str(info.meta.get("preprocess_meta")),
            str(info.meta.get("engine_options")),
        ])

    # Manifest sheet
    ws_m = wb.create_sheet(title="Manifest")
    for r, row_vals in enumerate(manifest_rows, start=1):
        for c, v in enumerate(row_vals, start=1):
            ws_m.cell(row=r, column=c, value=v)

    wb.save(output_path)